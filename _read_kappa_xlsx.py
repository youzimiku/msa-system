# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
p = r'C:\Users\youzi\xwechat_files\wxid_y2a1nw2f8c9u22_67ae\msg\file\2026-09\MCU6焊接KAPPA报告1.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
print('SHEETS:', wb.sheetnames)
for ws in wb.worksheets:
    print('=== sheet:', ws.title, 'dims:', ws.dimensions, 'max_row:', ws.max_row, 'max_col:', ws.max_column)
    # 合并单元格
    if ws.merged_cells.ranges:
        print('merged:', [str(r) for r in ws.merged_cells.ranges][:40])
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80)):
        vals = []
        for c in row:
            v = c.value
            if v is not None:
                vals.append('%s%s=%r' % (c.coordinate, '(M)' if c.data_type == 's' else '', v))
        if vals:
            print(' | '.join(vals))
