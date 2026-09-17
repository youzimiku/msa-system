# -*- coding: utf-8 -*-
"""生成 MSA 系统功能清单 xlsx（本地 Excel）"""
import io, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r'C:\Users\youzi\Doubao\chats\2026-09-03\new-chat\MSA系统\MSA系统功能清单.xlsx'

wb = Workbook()

# ---------- 样式 ----------
HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
BODY_FONT = Font(name='微软雅黑', size=10, color='333333')
BOLD_FONT = Font(name='微软雅黑', size=10, bold=True, color='333333')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical='center', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

STATUS_FILL = {
    '已有': PatternFill('solid', fgColor='E8F5E9'),   # 绿
    '新增': PatternFill('solid', fgColor='FFEBEE'),   # 红
    '调整': PatternFill('solid', fgColor='E3F2FD'),   # 蓝
    '重构': PatternFill('solid', fgColor='FFF3E0'),   # 橙
    '待确认': PatternFill('solid', fgColor='F3E5F5'), # 紫
}
PRIO_FILL = {
    'P0': PatternFill('solid', fgColor='FFCDD2'),
    'P1': PatternFill('solid', fgColor='FFE0B2'),
    'P2': PatternFill('solid', fgColor='FFF9C4'),
}

def style_sheet(ws, headers, rows, widths, status_col=None, prio_col=None, freeze='A2'):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    for r in rows:
        ws.append(r)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER
        if status_col:
            sc = ws.cell(row=r, column=status_col)
            v = str(sc.value or '')
            if v in STATUS_FILL:
                sc.fill = STATUS_FILL[v]
                sc.font = BOLD_FONT
                sc.alignment = CENTER
        if prio_col:
            pc = ws.cell(row=r, column=prio_col)
            v = str(pc.value or '')
            if v in PRIO_FILL:
                pc.fill = PRIO_FILL[v]
                pc.font = BOLD_FONT
                pc.alignment = CENTER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)), len(rows) + 1)
    ws.row_dimensions[1].height = 22

# ---------- Sheet1 功能清单主表 ----------
ws1 = wb.active
ws1.title = '功能清单'
h1 = ['序号', '模块', '功能点', '业务说明', '当前状态', '优先级', '核对点 / 备注']
rows1 = [
    [1, 'M1 质量特性维护', '特性定义', '特性编号/名称、关联量具、默认分析方法、默认抽样规则（驱动计划自动带出，对应「特性控制计划」层）', '新增', 'P1', ''],
    [2, 'M1 质量特性维护', '特性增删改查', '特性列表 + 编辑 + 启用/停用', '新增', 'P1', ''],
    [3, 'M1 质量特性维护', '特性与抽样方法关系', '质量特性与抽样方法是合并一张表维护还是分两张表维护', '待确认', 'P0', '与石洋洋确认'],
    [4, 'M2 抽样规则/抽样方法维护', '抽样方法定义', '方法名称、所属检查方法（抽样方法≠检验标准，检验标准绑零件，抽样方法绑检查方法）', '新增', 'P1', '会议点名缺失'],
    [5, 'M2 抽样规则/抽样方法维护', '取样数量规则', '样品数/人数/次数：GRR 10×3×3、KAPPA 50×3×3、线性 5标准件×12、稳定性 25子组×5、CgCgk 50次；界面按默认设置、可调', '新增', 'P1', '业务确认的取样规则'],
    [6, 'M2 抽样规则/抽样方法维护', '是否走真值', '标注该抽样方法是否依赖真值分析（联动样本管理「真值」字段）', '新增', 'P2', ''],
    [7, 'M2 抽样规则/抽样方法维护', '合并取样关系', '哪些抽样方法需要合并取样，在技术数据中维护', '新增', 'P2', ''],
    [8, 'M3 计量器具台账', '器具增删改查', '器具编号/名称/型号/量程/精度/状态等', '已有', 'P1', ''],
    [9, 'M3 计量器具台账', '校准管理', '登记校准、到期提醒、不合格自动停用、状态回写台账', '已有', 'P1', ''],
    [10, 'M3 计量器具台账', '器具组逻辑', '先有器具台账、再选编组；无组器具照常可查可用；器具组靠数据导入。当前是先建组再导器具、创建计划强依赖选组，逻辑反了', '调整', 'P1', '会议明确反转'],
    [11, 'M3 计量器具台账', '样机标识', '台账中标记「样机」（系统已有 prototype 字段，确认展示方式）', '已有', 'P2', ''],
    [12, 'M4 样本管理', '样本组维护', '样本组编号/名称/关联计划/样本数/覆盖变差/状态', '已有', 'P1', ''],
    [13, 'M4 样本管理', '样本明细', '样本编号/参考值/测量状态', '已有', 'P1', ''],
    [14, 'M4 样本管理', '真值字段', '样本中增加「真值」录入（值 + 来源 + 测定设备），供走真值的抽样分析使用', '新增', 'P1', '会议点名'],
    [15, 'M5 MSA 计划管理', '特性控制计划', '特性层：定义每个特性下用哪些方法、抽样规则、分阶段安排（样品阶段只做CgK → 试用阶段加GRR/KAPPA → 量产/周期检验做全项）', '新增', 'P1', '三层架构第一层'],
    [16, 'M5 MSA 计划管理', 'MSA 总计划', '按量具建立：关联零件 + 测量哪一项数据 + 分析方法/抽样规则按量具特性默认带出（可改）；作为父级汇总所有子计划结论，主计划有汇总结论', '重构', 'P1', '三层架构第二层'],
    [17, 'M5 MSA 计划管理', 'MSA 分计划（取样计划）', '每个分析方法一个独立取样计划：独立抽样数据 + 独立结论；一个 MSA 总计划下分多个子计划', '新增', 'P1', '三层架构第三层'],
    [18, 'M5 MSA 计划管理', '计划列表展示', '平铺/上下结构直观显示「7个方法里用了哪几个、每个方法用的什么方法、结论是什么」，不能点开才看', '调整', 'P1', '会议点名展示形式'],
    [19, 'M6 分析台账与录入', '5 套分析台账', 'GRR / KAPPA / 线性偏移 / 稳定性 / CgCgk 各自独立台账页面', '已有', 'P1', ''],
    [20, 'M6 分析台账与录入', '台账字段差异化', '各方法台账字段应不同，哪些字段必须显示需与业务确认（当前 5 套字段相同，不对）', '调整', 'P1', '与石洋洋确认字段清单'],
    [21, 'M6 分析台账与录入', '5 套录入数据页', 'GRR/KAPPA 沿用既有录入布局，线性/稳定性/CgCgk 按同类样式；取样参数（人数/次数/样本数）可调', '已有', 'P1', '会议认可布局'],
    [22, 'M7 分析结果与报告', '结果计算与判定', '各方法 KPI 计算 + 判定准则（AIAG/业务三档准则已实现：NDC≥5、GRR% 阈值、Kappa>0.75、Cg/Cgk≥1.33、线性偏倚四准则等）', '已有', 'P1', ''],
    [23, 'M7 分析结果与报告', '报告展示样式', '数据展示样式（交叉表 vs DSO 直出 vs 特殊格式）、行/列定义、数据值填充、字段明细，需业务方给样例后开发', '待确认', 'P0', '与石洋洋确认；报告模块后置'],
    [24, 'M8 系统支撑', '基础支撑', '仪表盘、角色权限、用户中心、检验标准维护（绑零件/工序，保留，与抽样方法区分）', '已有', 'P2', ''],
]
style_sheet(ws1, h1, rows1, [6, 22, 20, 60, 10, 8, 22], status_col=5, prio_col=6)

# ---------- Sheet2 模块概览 ----------
ws2 = wb.create_sheet('模块概览')
h2 = ['模块编号', '模块名称', '整体状态', '模块用途 / 说明']
rows2 = [
    ['M1', '质量特性维护', '新增', '特性层基础数据：定义量具测量特性，驱动 MSA 计划自动带出方法与抽样规则'],
    ['M2', '抽样规则/抽样方法维护', '新增', '抽样方法、取样数量规则、是否走真值、合并取样关系维护（本次会议点名缺失）'],
    ['M3', '计量器具台账', '已有·调整', '器具维护 + 校准管理 + 器具组（逻辑反转：先有器具再编组）+ 样机标识'],
    ['M4', '样本管理', '已有·调整', '样本组/样本明细维护，新增「真值」字段支持走真值的分析'],
    ['M5', 'MSA 计划管理', '重构', '三层架构：特性控制计划 → MSA 总计划 → MSA 分计划（取样计划）；分阶段执行'],
    ['M6', '分析台账与录入', '已有·调整', '5 套独立台账 + 5 套独立录入数据页；台账字段按方法差异化'],
    ['M7', '分析结果与报告', '已有·待定', '结果计算与判定已实现；报告数据展示样式待业务给样例后开发'],
    ['M8', '系统支撑', '已有', '仪表盘、角色权限、用户中心、检验标准维护'],
]
style_sheet(ws2, h2, rows2, [10, 24, 16, 60])

# ---------- Sheet3 待确认事项 ----------
ws3 = wb.create_sheet('待确认事项')
h3 = ['序号', '待确认问题', '影响范围', '建议确认对象', '备注']
rows3 = [
    [1, '质量特性与抽样方法是合并一张表维护还是分两张表维护', 'M1/M2 数据结构', '石洋洋', '会议纪要待办，影响建表'],
    [2, 'MSA 计划三层架构（特性控制计划/MSA总计划/MSA分计划）是否可合并、具体业务逻辑', 'M5 整体重构', '石洋洋', '张会娟也表示拿不准，需与业务/用户核对'],
    [3, '各分析报告（GRR/CgK 等）的数据展示样式：交叉表还是 DSO 直出、行/列定义、数据值填充、字段明细', 'M7 报告模块', '石洋洋', '会议明确：报告最复杂，先等样式再开发'],
    [4, '各方法台账必显示字段清单（哪些字段必须、哪些不需要显示）', 'M6 台账字段', '石洋洋', '当前 5 套字段相同，需差异化'],
    [5, '器具组数据导入方式与无组器具的独立操作确认', 'M3 器具组逻辑', '石洋洋', '会议明确：先有器具再编组，无组不影响操作'],
    [6, '报告模块实施顺序（先做基础数据+计划，报告后置）', '实施排期', '高国翔/张会娟', '会议确定分步实施'],
]
style_sheet(ws3, h3, rows3, [6, 56, 24, 16, 34])

wb.save(OUT)
print('已生成:', OUT)
print('大小:', os.path.getsize(OUT), 'bytes')
