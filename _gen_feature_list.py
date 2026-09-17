# -*- coding: utf-8 -*-
"""根据当前 MSA 系统（index.html v2.6+层级+双表）生成功能清单本地 Excel"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows = [
    # (功能模块, 功能点, 功能描述)
    ("量具管理", "计量器具台账", "上下双列表布局：上=器具组列表（组编号/名称/成员数量/样机/备注/维护人），下=组内器具明细；器具组支持新增/编辑/删除；器具支持新增/编辑/查看/登记校准/状态变更/删除；查询条件含编号/名称/器具类别/状态/到期预警/未来N天；支持导入/导出。"),
    ("量具管理", "样机属性", "器具可标记为「样机」（原默认器具概念已统一为样机属性，界面显示样机列）。"),
    ("量具管理", "校准管理", "校准记录查看/登记/编辑，与器具关联；校准状态（正常/超期/临期）展示与到期预警；校准管理去 Tab 化、支持校准过期状态筛选。"),
    ("MSA管理", "MSA计划-查询筛选", "查询条件：计划号/关键词、分析关联类型（未定型/GRR/KAPPA/多方法）、计划状态（待开始/进行中/已完成）；查询/重置按钮置操作区最前。"),
    ("MSA管理", "MSA计划-创建", "创建弹窗：器具筛选（复评状态+厂商/型号/产线/工序）置顶；计划填写信息单块 6 列自适应（零件/工序、检验标准、质量特性、创建方式、计划名称模板、触发依据、计划检期、部门、零件号、质检区划、工厂、分厂、观察员、复评周期、复评提前提醒、测量人员多选）；分析方法多选；选择器具组并勾选器具；已存在未闭环计划的器具自动跳过。"),
    ("MSA管理", "MSA计划-取样规则", "按抽样方法维护的默认取样规则（GRR 10件×3人×3次、KAPPA 50件×3人×3次、线性 5标准件×12次、稳定性 25子组×5次、Cg/Cgk 50次），每行人数/次数/样本数可调节；勾选方法后显示分计划预览条（将生成 N 个分计划及对应取样规则）。"),
    ("MSA管理", "MSA计划-上下双表", "上表=MSA 计划列表（操作列：台账/编辑/详情，点击行选中并高亮）；下表=当前选中计划的分析单信息（分计划号/分析方法/取样规则/状态/结论/操作：查看结果/台账），默认选中第一条计划并联动切换。"),
    ("MSA管理", "MSA计划-详情", "计划详情抽屉：量具信息、计划信息、检验标准、判定结果、执行信息、备注；分计划列表面板（取样计划→录入→分析→结论归集，任一不合格则总计划不合格）；未定型计划可在详情内「转分析方法」定型并生成台账待采集记录。"),
    ("MSA管理", "MSA计划-复评机制", "复评周期（月）+复评提前提醒（天）配置；器具复评状态筛选与预警（正常/临期/超期）；周期复评需上一计划闭环后重新发起。"),
    ("MSA管理", "检验标准维护", "检验标准增删改查：编号/名称/零件/工序/检验项目/检验方法/数据类型/判定准则；标准绑定零件/工序/质量特性（不直接绑器具）；按分析方法类型（GRR/KAPPA/线性/稳定性/Cg/Cgk）匹配可用标准。"),
    ("MSA管理", "质量特性维护", "质量特性增删改查：特性编号/名称/关联零件/工序/特性类型（SC/CC/普通 三级）/规格上下限/单位；特性与检验标准、MSA 计划关联。"),
    ("MSA管理", "抽样方法维护", "7 大分析方法代码表（GRR/KAPPA/线性/偏倚/稳定性/分辨力/Cg/Cgk：方法编号/名称/数据类型/适用场景）；取样规则表（各方法默认 样本数/人数/次数，可调节）；分辨力不取样直接录入；线性+偏倚可合并取样。"),
    ("MSA管理", "样本管理", "样本组列表+样本明细抽屉：样本组编号/零件/检验标准/数量；样本组新增/编辑/删除；样本明细查看；样本与分计划关联。"),
    ("台账", "GRR台账", "GRR 台账记录列表（分析单号/计划号/器具/状态/结论/%GRR/NDC）；「录入数据」入口跳转录入页；详情/审核操作。"),
    ("台账", "KAPPA台账", "KAPPA 台账记录列表（分析单号/总体 KAPPA/状态/结论）；录入入口；详情/审核操作。"),
    ("台账", "线性/偏移台账", "线性/偏移台账记录列表（分析单号/回归斜率/截距/偏倚/状态/结论）；录入入口；详情/审核操作。"),
    ("台账", "稳定性台账", "稳定性台账记录列表（分析单号/状态/结论）；录入入口；详情/审核操作。"),
    ("台账", "Cg/Cgk台账", "Cg/Cgk 台账记录列表（分析单号/Cg/Cgk/状态/结论）；录入入口；详情/审核操作。"),
    ("录入数据", "GRR录入数据", "操作员×样本×试验矩阵录入（行=操作员、列=样本×次数），按取样规则（10件×3人×3次可调）校验完整性；提交后自动计算 GRR 指标并置为「待审核」。"),
    ("录入数据", "KAPPA录入数据", "判定矩阵录入（检验员×样本 合格/不合格 0/1），按取样规则（50件×3人×3次可调）校验；自动计算 KAPPA 与有效性/错误率/错误报警率。"),
    ("录入数据", "线性/偏移录入数据", "标准件×次数测量值表录入（5标准件×12次可调）；自动计算线性回归（斜率/截距）与各点偏倚。"),
    ("录入数据", "稳定性录入数据", "子组×次数表录入（25子组×5次可调，行可增删）；自动计算子组均值/极差等稳定性控制指标。"),
    ("录入数据", "Cg/Cgk录入数据", "连续测量次数表录入（50次可调）；自动计算 Cg/Cgk 指数与重复性误差占比 6σ/T。"),
    ("分析执行", "GRR分析结果", "KPI 展示：%GRR、NDC、EV 设备变差、AV 操作员变差等；判定准则（NDC≥5 接受；%GRR<10% 可接受、10%~30% 勉强接受、>30% 不可接受）；结论状态（待审核/已批准/退回整改/整改闭环）；审核通过/退回+纠正措施+整改闭环联动计划结论。"),
    ("分析执行", "KAPPA分析结果", "总体 KAPPA 展示；判定准则（KAPPA≥0.75 满足要求、0.40~0.75 边缘、<0.40 不可接受）；有效性/错误率/错误报警率；报告表单号 GJZ-MSA-xxx；审核与整改闭环。"),
    ("分析执行", "线性/偏移分析结果", "线性回归方程（斜率/截距）、各点平均偏倚与置信区间；4 档判定准则（零偏倚水平线完全在置信区间内=非常理想可接受；固定偏倚可纠偏修正=理想可接受；线性偏倚可回归修正=较理想可接受；无线性不可修正=不可接受）；审核与整改闭环。"),
    ("分析执行", "稳定性分析结果", "子组均值/极差等稳定性指标展示；稳定性判定；审核与整改闭环。"),
    ("分析执行", "Cg/Cgk分析结果", "Cg/Cgk 指数展示；判定准则（均≥1.33 量具能力充足；Cg 偏小重复性差需检修；Cg 合格 Cgk 偏小存在系统偏倚需校准；Cgk 为负偏倚超 10% 公差不可使用）；重复性误差占比 6σ/T（≤15% 优秀、≤20% 可接受）；审核与整改闭环。"),
    ("分析执行", "分辨率分析", "分辨力值直接录入（不取样）并判定，审核与整改闭环。"),
    ("通用能力", "权限控制", "按角色（质量工程师等）控制新增/编辑/审核等操作权限（canDo 校验）。"),
    ("通用能力", "数据持久化", "演示数据存 localStorage（msa_demo_data），浏览器本地持久化。"),
    ("通用能力", "统一 UI 规范", "Ant Design 组件库；页面统一「查询条件/操作/列表」三面板；一列一字段；输入控件统一宽度；状态/结论字段宽度统一并支持鼠标悬停说明；说明性文字已按要求移除。"),
]

wb = Workbook()
ws = wb.active
ws.title = "功能清单"

headers = ["功能模块", "功能点", "功能描述"]
ws.append(headers)

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F4E79")
head_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
body_font = Font(name="微软雅黑", size=10, color="333333")
module_fill = PatternFill("solid", fgColor="DCE6F1")

for c in range(1, 4):
    cell = ws.cell(row=1, column=c)
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

r = 2
for mod, pt, desc in rows:
    ws.append([mod, pt, desc])
    ws.cell(row=r, column=1).font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
    ws.cell(row=r, column=1).fill = module_fill
    ws.cell(row=r, column=1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(row=r, column=2).font = Font(name="微软雅黑", size=10, bold=True)
    ws.cell(row=r, column=2).alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(row=r, column=3).font = body_font
    ws.cell(row=r, column=3).alignment = Alignment(vertical="center", wrap_text=True)
    for c in range(1, 4):
        ws.cell(row=r, column=c).border = border
    r += 1

ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 95
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 24

out = r"C:\Users\youzi\Doubao\chats\2026-09-03\new-chat\MSA系统\MSA系统功能清单（当前版）.xlsx"
wb.save(out)
print("saved", out, "rows", len(rows))
