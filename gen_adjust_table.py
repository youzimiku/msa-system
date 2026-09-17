# -*- coding: utf-8 -*-
"""生成《MSA系统调整项处理表.xlsx》：方案有、用户清单未记的调整项，供用户填写处理意见"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

rows = [
    ("序号", "模块", "调整项", "说明 / 背景", "当前状态", "如何处理（待填写）"),
    (1, "MSA计划创建", "新增字段：零件号、被测特性",
     "会议决议7：计划创建必须包含零件号、被测特性、特性级别、分析人。特性级别你已确认暂不管，零件号与被测特性是否补上？",
     "未实施"),
    (2, "MSA计划创建", "器具固有属性完整剥离（复评周期、提前提醒回归台账维护）",
     "决议2：分辨率、精度、复评周期、提前提醒天数均为器具固有属性，应由计量器具台账维护、计划侧自动带出。目前：分辨力已改文本输入、精度已加到台账、创建弹窗复评两个输入框已删，但台账侧尚无复评周期/提前提醒维护字段。",
     "部分实施"),
    (3, "MSA计划创建", "分析人=计划总负责人（任务分配/权限控制）",
     "决议4：分析人不只是名字，应作为该计划总负责人，用于后续任务分配与权限控制。目前仅为列表字段。",
     "未实施"),
    (4, "MSA计划", "一量具一计划、一主计划多方法多子任务",
     "决议5：选择多个量具生成多个主计划；一个主计划勾选多个分析方法生成多个子任务（台账），父子可关联。你之前专门问过该逻辑。",
     "未实施"),
    (5, "MSA计划层级", "详情卡片化展示子任务及结论",
     "决议6：弃用上下双列表，主计划列表+点详情以卡片形式查看各子任务及各子任务结论。目前去掉双列表后此替代交互未做。",
     "未实施"),
    (6, "检验标准/取样", "取样规则三级继承",
     "决议11：一级分析方法维护默认规则（GRR 10件×3人×3次等）→二级检验标准带出可调→三级计划创建带出可微调。目前取样规则仅存于台账记录，三级继承链路未做。",
     "未实施"),
    (7, "数据录入", "录入矩阵动态生成",
     "决议14：按取样规则自动生成录入矩阵（如10件×3人×3次=90行），支持加减号动态增删行、调整参数后重新生成。目前为固定矩阵。",
     "未实施"),
    (8, "数据录入", "录入界面隐藏干扰信息",
     "决议15：录入界面只显示矩阵及基础信息（零件号/零件名/工序/量具），严禁显示真值、上下限、分辨率。",
     "未实施"),
    (9, "台账", "导入数据预览子页",
     "决议20：导入Excel后先进入预览页确认数据无误，再手动执行分析。目前导入即生效。",
     "未实施"),
    (10, "台账", "样本件不进台账列表",
     "决议17后半：样本件不必在台账中列出，随录入记录体现。目前台账记录层面未涉及样本件展示，基本符合，待确认。",
     "部分实施"),
    (11, "分析结果", "动态图表为主 + 悬停查看数据点；静态报告暂缓",
     "决议18：分析结果以动态图表呈现（悬停显示数据点信息），一页纸静态报告暂缓至9月底再议。目前为静态结论+KPI。",
     "未实施"),
    (12, "权限", "权限分级：操作者仅录入页，分析人全量",
     "决议17：操作者（操作人）仅拥有数据录入权限，不开放台账与真值；分析人拥有全量权限（台账、数据、分析结果）。",
     "未实施"),
    (13, "质量特性", "特性结构化（过程/产品特性+等级代码、计量/技术型、公差拆三字段、引用日志穿透CP）",
     "决议9：参照SPC52特殊特性清单，特性类型区分过程/产品特性并带等级代码、特性类别分计量/技术型、公差拆标准值/上限/下限、支持引用日志与CP穿透追溯。特性类型SC2你已确认不加，特性级别代码未定。",
     "暂缓（特性类型不加SC2；级别代码未定）"),
    (14, "检验标准", "与老系统检验标准字段合并",
     "决议12：与老系统（项目号/项目名/文件号/工序/方法/取样规则/分辨率）取最大值合并。你已确认不做（当前维护字段更全面）。",
     "已确认不做"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "调整项处理表"

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="E8E8E8")
head_font = Font(bold=True, size=11, color="333333")
body_font = Font(size=11, color="333333")
wrap = Alignment(vertical="center", wrap_text=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

widths = [6, 16, 34, 60, 18, 20]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + i)].width = w

for c in range(1, 7):
    cell = ws.cell(row=1, column=c)
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = center
    cell.border = border

for r in rows:
    ws.append(r)

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
    for cell in row:
        cell.font = body_font
        cell.border = border
        cell.alignment = center if cell.column in (1,) else wrap

ws.row_dimensions[1].height = 24
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 46

out = r"C:\Users\youzi\Doubao\chats\2026-09-03\new-chat\MSA系统\MSA系统调整项处理表.xlsx"
wb.save(out)
print("saved:", out, "rows:", ws.max_row)
