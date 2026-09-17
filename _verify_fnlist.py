# -*- coding: utf-8 -*-
"""回读校验功能清单 xlsx"""
import io
from openpyxl import load_workbook

P = r'C:\Users\youzi\Doubao\chats\2026-09-03\new-chat\MSA系统\MSA系统功能清单.xlsx'
wb = load_workbook(P)
print('工作表:', wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print('--', name, '行数:', ws.max_row, '列数:', ws.max_column)
    # 打印表头
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print('  表头:', hdr)
    # 抽查末行
    last = [ws.cell(row=ws.max_row, column=c).value for c in range(1, ws.max_column + 1)]
    print('  末行:', [str(x)[:30] if x else '' for x in last])
