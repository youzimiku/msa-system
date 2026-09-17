# -*- coding: utf-8 -*-
import io, json, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows = json.load(io.open(r'C:\Users\youzi\Doubao\chats\2026-09-03\new-chat\MSA系统\_pageguide_rows.json', encoding='utf-8'))
out = r'C:\Users\youzi\Doubao\chats\2026-09-03\new-chat\MSA系统\MSA系统_页面说明_20260917.xlsx'
wb = Workbook()
ws = wb.active
ws.title = '页面说明'
headers = ['页面Key', '页面名称', '段落', '段落键', '标识（仅状态行）', '当前文字']
ws.append(headers)
thin = Side(style='thin', color='d9d9d9')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill('solid', fgColor='FAFAFA')
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = hfill
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = border
for r in rows:
    ws.append(r)
widths = [11, 15, 13, 11, 20, 70]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6):
    for c in row:
        c.border = border
        c.alignment = Alignment(vertical='top', wrap_text=True)
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:F%d' % ws.max_row
wb.save(out)
print('已生成:', out, '行数:', ws.max_row - 1)
